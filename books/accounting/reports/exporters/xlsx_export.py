"""
XLSX exporter for the trial balance.

Same shape as the CSV (Q2 refinement — no type subheader rows;
`type` column does the grouping; depth column + indented name for
hierarchy). Differs from CSV in:

  - Decimal cells use Excel native numbers, not strings, so Excel
    can SUM / pivot / format them.
  - Number format `#,##0.00;(#,##0.00)` mirrors HTML's parens-for-
    negatives convention.
  - Header row + totals row are bold.
  - Indent uses leading spaces (4 per depth level) inside the name
    cell. Spaces survive copy/paste into other tools, unlike
    openpyxl's alignment.indent.

The cell-level round-trip test reads `cell.value` directly (not the
formatted display) so the contract is the underlying Decimal, not
the display string.
"""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from books.accounting.reports.trial_balance import TrialBalance, TrialBalanceRow


COLUMNS_BASE = [
    "account_number",
    "name",
    "type",
    "depth",
    "debits",
    "credits",
    "own_balance",
    "rollup_balance",
]
PRIOR_COLUMN = "prior_own_balance"
INDENT = "    "  # 4 spaces per depth level
NUMERIC_COLUMNS = {"debits", "credits", "own_balance", "rollup_balance",
                   "prior_own_balance"}
ACCOUNTING_FORMAT = '#,##0.00;(#,##0.00)'


def render_xlsx(tb: TrialBalance) -> bytes:
    """Render the trial balance as an XLSX byte string."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Trial Balance"

    columns = list(COLUMNS_BASE)
    if tb.prior_as_of is not None:
        columns.append(PRIOR_COLUMN)

    bold = Font(bold=True)

    # Header row.
    ws.append(columns)
    for cell in ws[1]:
        cell.font = bold

    # Data rows.
    for row in _walk(tb.rows):
        ws.append(_row_values(row, tb.prior_as_of is not None))
        # Apply number format on numeric columns of this row.
        excel_row = ws.max_row
        for idx, col in enumerate(columns, start=1):
            if col in NUMERIC_COLUMNS:
                ws.cell(row=excel_row, column=idx).number_format = ACCOUNTING_FORMAT

    # Totals row (bold).
    totals = ["TOTAL", "", "", "", tb.total_debits, tb.total_credits,
              "Balanced" if tb.is_balanced else "OUT OF BALANCE", ""]
    if tb.prior_as_of is not None:
        totals.append("")
    ws.append(totals)
    excel_row = ws.max_row
    for cell in ws[excel_row]:
        cell.font = bold
    # Numeric format on the totals' debits/credits columns (5 and 6).
    ws.cell(row=excel_row, column=5).number_format = ACCOUNTING_FORMAT
    ws.cell(row=excel_row, column=6).number_format = ACCOUNTING_FORMAT

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row_values(row: TrialBalanceRow, with_prior: bool) -> list:
    values = [
        row.account.account_number,
        INDENT * (row.depth - 1) + row.account.name,
        row.account.type,
        row.depth,
        row.debits_total,
        row.credits_total,
        row.own_balance,
        row.rollup_balance,
    ]
    if with_prior:
        values.append(
            row.prior_own_balance if row.prior_own_balance is not None else None
        )
    return values


def _walk(rows):
    for row in rows:
        yield row
        yield from _walk(row.children)
